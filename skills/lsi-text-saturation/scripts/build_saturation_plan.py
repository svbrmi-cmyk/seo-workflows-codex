#!/usr/bin/env python3
"""Merge median and full-depth workbooks into a saturation audit."""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from validate_lsi_saturation import (
    anti_monoculture,
    as_int,
    count_group,
    normalize,
    paragraph_count,
    local_repetitions,
    read_terms,
    tokens,
)


@dataclass
class TableRow:
    row: int
    label: str
    forms: tuple[str, ...]
    analyzer_count: int
    median_seen: bool = False
    median: int = 0
    minimum: int = 0
    maximum: int = 0
    recommendation: int = 0


@dataclass
class PlanGroup:
    row: int
    label: str
    forms: tuple[str, ...]
    analyzer_count: int
    median_seen: bool
    median: int
    minimum: int
    maximum: int
    recommendation: int

    @property
    def target(self) -> int:
        if self.recommendation < 0:
            return 0
        if self.median_seen:
            return self.median
        return self.minimum

    @property
    def cap(self) -> int:
        if self.recommendation < 0:
            return self.maximum if self.maximum > 0 else 0
        if self.median_seen:
            return self.median
        if self.maximum > 0:
            return self.maximum
        return self.target

    @property
    def negative(self) -> bool:
        return self.recommendation < 0


def headers(row: Iterable[object]) -> dict[str, int]:
    return {normalize(value): index for index, value in enumerate(row) if value is not None}


def column(mapping: dict[str, int], *names: str) -> int | None:
    for name in names:
        key = normalize(name)
        if key in mapping:
            return mapping[key]
    return None


def cell(row: tuple[object, ...], index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def split_forms(value: object, label: str) -> tuple[str, ...]:
    forms = {label}
    for item in str(value or "").replace(";", ",").split(","):
        form = normalize(item)
        if form:
            forms.add(form)
    return tuple(sorted(forms))


def read_table(path: Path, median_mode: bool) -> list[TableRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    selected = None
    mapping: dict[str, int] = {}
    for sheet in workbook.worksheets:
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        candidate = headers(first)
        if column(candidate, "Слово", "Слова") is None:
            continue
        required = column(candidate, "Медиана") if median_mode else column(candidate, "Минимум по рекомендациям")
        if required is not None:
            selected, mapping = sheet, candidate
            break
    if selected is None:
        kind = "медианы" if median_mode else "полной глубины"
        raise ValueError(f"Не найден лист {kind} в {path}")

    word_col = column(mapping, "Слово", "Слова")
    forms_col = column(mapping, "Словоформы")
    count_col = column(mapping, "Повторы", "Текст у вас")
    median_col = column(mapping, "Медиана")
    minimum_col = column(mapping, "Минимум по рекомендациям")
    maximum_col = column(mapping, "Максимум по рекомендациям")
    recommendation_col = column(mapping, "Добавить/Удалить")
    assert word_col is not None

    result: list[TableRow] = []
    for number, row in enumerate(selected.iter_rows(min_row=2, values_only=True), start=2):
        label = normalize(cell(row, word_col))
        if not label:
            continue
        result.append(TableRow(
            row=number,
            label=label,
            forms=split_forms(cell(row, forms_col), label),
            analyzer_count=as_int(cell(row, count_col)),
            median_seen=median_mode,
            median=as_int(cell(row, median_col)),
            minimum=as_int(cell(row, minimum_col)),
            maximum=as_int(cell(row, maximum_col)),
            recommendation=int(round(float(str(cell(row, recommendation_col) or 0).replace(",", ".")))),
        ))
    return result


def find_match(row: TableRow, groups: list[PlanGroup]) -> PlanGroup | None:
    for group in groups:
        if group.label == row.label:
            return group
    row_forms = set(row.forms)
    matches = [group for group in groups if row_forms.intersection(group.forms)]
    return matches[0] if len(matches) == 1 else None


def merge_tables(median_rows: list[TableRow], deep_rows: list[TableRow]) -> tuple[list[PlanGroup], list[str]]:
    groups = [
        PlanGroup(
            row=row.row,
            label=row.label,
            forms=row.forms,
            analyzer_count=row.analyzer_count,
            median_seen=False,
            median=0,
            minimum=row.minimum,
            maximum=row.maximum,
            recommendation=row.recommendation,
        )
        for row in deep_rows
    ]
    warnings: list[str] = []

    for row in median_rows:
        match = find_match(row, groups)
        if match is None:
            groups.append(PlanGroup(
                row=row.row,
                label=row.label,
                forms=row.forms,
                analyzer_count=row.analyzer_count,
                median_seen=True,
                median=row.median,
                minimum=0,
                maximum=0,
                recommendation=row.recommendation,
            ))
            warnings.append(f"Группа медианы отсутствует в полной таблице: {row.label}")
            continue
        match.forms = tuple(sorted(set(match.forms).union(row.forms)))
        match.median_seen = True
        match.median = row.median
        if row.analyzer_count != match.analyzer_count:
            warnings.append(
                f"Разные счетчики таблиц для {match.label}: "
                f"{row.analyzer_count} и {match.analyzer_count}"
            )
    return groups, warnings


def state(group: PlanGroup, exclude: set[str]) -> str:
    names = {group.label, *group.forms}
    if not names.isdisjoint(exclude):
        return "excluded"
    if group.negative:
        return "negative"
    if group.median_seen and group.median == 0:
        return "median-zero"
    if group.target > 0:
        return "positive"
    return "ignored"


def audit_plan(
    groups: list[PlanGroup],
    source: str,
    edited: str,
    exclude: set[str],
) -> dict[str, object]:
    before_tokens = tokens(source)
    after_tokens = tokens(edited)
    rows: list[dict[str, object]] = []

    for group in groups:
        group_state = state(group, exclude)
        before = count_group(before_tokens, group)
        after = count_group(after_tokens, group)
        target = group.target
        cap = group.cap
        rows.append({
            "row": group.row,
            "group": group.label,
            "forms": list(group.forms),
            "state": group_state,
            "analyzer_count": group.analyzer_count,
            "source_count": before,
            "edited_count": after,
            "count_mismatch": before != group.analyzer_count,
            "median_seen": group.median_seen,
            "median": group.median,
            "minimum": group.minimum,
            "maximum": group.maximum,
            "recommendation": group.recommendation,
            "target": target,
            "cap": cap,
            "covered_before": group_state == "positive" and before > 0,
            "covered_after": group_state == "positive" and after > 0,
            "target_met": group_state != "positive" or after >= target,
            "over_cap": group_state in {"positive", "negative"} and after > cap,
            "negative_met": group_state != "negative" or after <= cap,
            "paragraphs": paragraph_count(edited, group),
            "clustered": group_state == "positive" and after >= 2 and paragraph_count(edited, group) < 2,
        })

    positive = [row for row in rows if row["state"] == "positive"]
    before_covered = sum(bool(row["covered_before"]) for row in positive)
    after_covered = sum(bool(row["covered_after"]) for row in positive)
    total_target = sum(int(row["target"]) for row in positive)
    before_depth = sum(min(int(row["source_count"]), int(row["target"])) for row in positive)
    after_depth = sum(min(int(row["edited_count"]), int(row["target"])) for row in positive)
    target_forms = {
        form
        for group in groups
        if state(group, exclude) == "positive"
        for form in group.forms
    }
    monoculture = anti_monoculture(source, edited, target_forms)
    active_groups = [group for group in groups if state(group, exclude) in {"positive", "negative"}]
    local = local_repetitions(edited, active_groups)

    return {
        "summary": {
            "groups": len(groups),
            "positive_groups": len(positive),
            "negative_groups": sum(row["state"] == "negative" for row in rows),
            "excluded_groups": sum(row["state"] == "excluded" for row in rows),
            "covered_before": before_covered,
            "covered_after": after_covered,
            "breadth_before": round(before_covered / len(positive), 4) if positive else 0.0,
            "breadth_after": round(after_covered / len(positive), 4) if positive else 0.0,
            "depth_before": round(before_depth / total_target, 4) if total_target else 0.0,
            "depth_after": round(after_depth / total_target, 4) if total_target else 0.0,
            "missing_width": sum(row["state"] == "positive" and int(row["edited_count"]) == 0 for row in rows),
            "target_deficits": sum(row["state"] == "positive" and not row["target_met"] for row in rows),
            "over_cap": sum(bool(row["over_cap"]) for row in rows),
            "negative_violations": sum(row["state"] == "negative" and not row["negative_met"] for row in rows),
            "clustered": sum(bool(row["clustered"]) for row in rows),
            "count_mismatches": sum(bool(row["count_mismatch"]) for row in rows),
            "monoculture_words": len(monoculture["words"]),
            "monoculture_ngrams": len(monoculture["ngrams"]),
            "same_sentence_repetitions": len(local["same_sentence_groups"]),
            "adjacent_sentence_repetitions": len(local["adjacent_groups"]),
            "paragraph_repetitions": len(local["paragraph_groups"]),
            "same_sentence_word_repetitions": len(local["same_sentence_words"]),
            "adjacent_sentence_word_repetitions": len(local["adjacent_words"]),
            "paragraph_word_repetitions": len(local["paragraph_words"]),
            "yo_symbols": edited.count("\u0451") + edited.count("\u0401"),
        },
        "groups": rows,
        "anti_monoculture": monoculture,
        "local_repetitions": local,
    }


def render(report: dict[str, object], warnings: list[str]) -> None:
    summary = report["summary"]
    print(
        "GROUPS={groups} POSITIVE={positive_groups} NEGATIVE={negative_groups} "
        "BREADTH={covered_before}->{covered_after} "
        "({breadth_before:.1%}->{breadth_after:.1%}) "
        "DEPTH={depth_before:.1%}->{depth_after:.1%}".format(**summary)
    )
    print(
        "MISSING_WIDTH={missing_width} TARGET_DEFICITS={target_deficits} "
        "OVER_CAP={over_cap} NEGATIVE_VIOLATIONS={negative_violations} "
        "CLUSTERED={clustered} MONOCULTURE={monoculture_words}/{monoculture_ngrams} "
        "LOCAL={same_sentence_repetitions}/{adjacent_sentence_repetitions}/{paragraph_repetitions} "
        "WORDS={same_sentence_word_repetitions}/{adjacent_sentence_word_repetitions}/{paragraph_word_repetitions} "
        "YO={yo_symbols}".format(**summary)
    )
    for warning in warnings:
        print(f"WARNING: {warning}")
    for row in report["groups"]:
        if row["state"] == "positive" and (
            int(row["edited_count"]) == 0 or not row["target_met"] or row["over_cap"]
        ):
            print(
                f"POSITIVE: {row['group']} {row['source_count']}->{row['edited_count']} "
                f"target={row['target']} cap={row['cap']}"
            )
        if row["state"] == "negative" and not row["negative_met"]:
            print(
                f"NEGATIVE: {row['group']} {row['source_count']}->{row['edited_count']} "
                f"cap={row['cap']}"
            )


def failures(report: dict[str, object]) -> list[str]:
    summary = report["summary"]
    checks = [
        ("есть непокрытые группы ширины", "missing_width"),
        ("есть недобор положительных целей", "target_deficits"),
        ("есть превышения лимитов", "over_cap"),
        ("не выполнены отрицательные рекомендации", "negative_violations"),
        ("есть скопления", "clustered"),
        ("целевая группа повторяется в одном предложении", "same_sentence_repetitions"),
        ("целевая группа повторяется в соседних предложениях", "adjacent_sentence_repetitions"),
        ("целевая группа чрезмерно сконцентрирована в одном абзаце", "paragraph_repetitions"),
        ("знаменательное слово повторяется в одном предложении", "same_sentence_word_repetitions"),
        ("знаменательное слово повторяется в соседних предложениях", "adjacent_sentence_word_repetitions"),
        ("знаменательное слово чрезмерно сконцентрировано в одном абзаце", "paragraph_word_repetitions"),
        ("найден символ U+0451", "yo_symbols"),
    ]
    return [message for message, key in checks if int(summary[key]) > 0]


def self_test() -> int:
    groups = [
        PlanGroup(2, "душевой", ("душевая", "душевой"), 0, True, 4, 1, 8, 4),
        PlanGroup(3, "смеситель", ("смеситель", "смесителем"), 0, True, 3, 1, 7, 3),
        PlanGroup(4, "керамика", ("керамика", "керамикой"), 0, False, 0, 1, 2, 1),
        PlanGroup(5, "конфиденциальность", ("конфиденциальность",), 0, False, 0, 10, 20, 10),
        PlanGroup(6, "удобный", ("удобный",), 3, False, 0, 0, 0, -3),
    ]
    source = "Удобный товар. Удобный выбор. Удобный вариант."
    edited = (
        "Душевая зона дополнена керамикой.\n\n"
        "Смеситель установлен рядом.\n\n"
        "Душевой узел работает стабильно.\n\n"
        "Для душевой ниши выбран смеситель.\n\n"
        "Душевой блок соединен со смесителем."
    )
    report = audit_plan(groups, source, edited, {"конфиденциальность"})
    summary = report["summary"]
    morphology = PlanGroup(
        7,
        "мебель",
        ("мебель", "мебели", "мебелью"),
        3,
        False,
        0,
        1,
        5,
        0,
    )
    morphology_count = count_group(
        tokens("Мебель дополняют светом. Для мебели важен уход. Между мебелью и стеной остается стык."),
        morphology,
    )
    negative_with_median = PlanGroup(
        8,
        "должно",
        ("должно", "должна", "должен", "должны"),
        4,
        True,
        5,
        0,
        2,
        -2,
    )
    checks = {
        "width_complete": summary["missing_width"] == 0,
        "targets_met": summary["target_deficits"] == 0,
        "caps_respected": summary["over_cap"] == 0,
        "negative_removed": summary["negative_violations"] == 0,
        "noise_excluded": summary["excluded_groups"] == 1,
        "distributed": summary["clustered"] == 0,
        "no_monoculture": summary["monoculture_words"] == 0 and summary["monoculture_ngrams"] == 0,
        "all_word_forms_counted": morphology_count == 3,
        "negative_overrides_median_target": negative_with_median.target == 0,
        "negative_uses_reduction_cap": negative_with_median.cap == 2,
    }
    for name, passed in checks.items():
        print(f"{'PASS' if passed else 'FAIL'} {name}")
    return 0 if all(checks.values()) else 1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path)
    parser.add_argument("--text", type=Path)
    parser.add_argument("--median", type=Path)
    parser.add_argument("--deep", type=Path)
    parser.add_argument("--exclude", type=Path)
    parser.add_argument("--json-output", type=Path)
    parser.add_argument("--strict", action="store_true")
    parser.add_argument("--self-test", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.self_test:
        return self_test()
    missing = [name for name in ("source", "text", "median", "deep") if getattr(args, name) is None]
    if missing:
        raise SystemExit("Обязательные параметры: --source, --text, --median, --deep")

    median_rows = read_table(args.median, median_mode=True)
    deep_rows = read_table(args.deep, median_mode=False)
    groups, warnings = merge_tables(median_rows, deep_rows)
    report = audit_plan(
        groups,
        args.source.read_text(encoding="utf-8-sig"),
        args.text.read_text(encoding="utf-8-sig"),
        read_terms(args.exclude),
    )
    render(report, warnings)
    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    issues = failures(report)
    if args.strict and issues:
        print("STRICT_FAIL: " + "; ".join(issues), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
