#!/usr/bin/env python3
"""Audit Russian LSI saturation against an XLSX median table."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Install openpyxl: python -m pip install openpyxl") from exc


TOKEN_RE = re.compile("[A-Za-zА-Яа-я\\u0401\\u04510-9]+(?:-[A-Za-zА-Яа-я\\u0401\\u04510-9]+)*")
PARAGRAPH_RE = re.compile(r"(?:\r?\n)\s*(?:\r?\n)+")
SENTENCE_RE = re.compile(r"(?<=[.!?])\s+")
STOPWORDS = {
    "а", "без", "более", "бы", "был", "была", "были", "было", "в", "вам",
    "вас", "весь", "во", "вот", "все", "для", "до", "его", "ее", "если",
    "есть", "еще", "же", "за", "и", "из", "или", "их", "к", "как", "ко",
    "который", "ли", "на", "над", "не", "но", "о", "об", "от", "по", "под",
    "перед", "после", "при", "с", "со", "так", "также", "то", "у", "уже", "что", "чтобы",
    "это", "этот",
}


@dataclass
class Group:
    row: int
    label: str
    median: int
    forms: tuple[str, ...]
    link: str


def normalize(value: object) -> str:
    return str(value or "").strip().lower().replace("\u0451", "е")


def as_int(value: object) -> int:
    if value in (None, ""):
        return 0
    try:
        return max(0, int(round(float(str(value).replace(",", ".")))))
    except (TypeError, ValueError):
        return 0


def tokens(text: str) -> list[str]:
    return [normalize(match.group(0)) for match in TOKEN_RE.finditer(text)]


def header_map(row: Iterable[object]) -> dict[str, int]:
    return {normalize(value): index for index, value in enumerate(row) if value is not None}


def find_column(headers: dict[str, int], *names: str) -> int | None:
    normalized = [normalize(name) for name in names]
    for name in normalized:
        if name in headers:
            return headers[name]
    return None


def read_groups(path: Path, limit: int) -> tuple[list[Group], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    warnings: list[str] = []
    main = None
    main_headers: dict[str, int] = {}

    for sheet in workbook.worksheets:
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = header_map(first)
        if find_column(headers, "Слова", "Слово") is not None and find_column(headers, "Медиана") is not None:
            if sheet.title == "Сводим формы слов":
                main, main_headers = sheet, headers
                break
            if main is None:
                main, main_headers = sheet, headers

    if main is None:
        raise ValueError("Не найден лист с колонками 'Слова/Слово' и 'Медиана'")

    word_col = find_column(main_headers, "Слова", "Слово")
    median_col = find_column(main_headers, "Медиана")
    link_col = find_column(main_headers, "Связь слов с другим листом")
    assert word_col is not None and median_col is not None

    raw: list[tuple[int, str, int, str]] = []
    seen: set[str] = set()
    for row_number, row in enumerate(main.iter_rows(min_row=2, values_only=True), start=2):
        label = normalize(row[word_col] if word_col < len(row) else "")
        if not label:
            continue
        link = normalize(row[link_col] if link_col is not None and link_col < len(row) else "")
        identity = link or label
        if identity in seen:
            continue
        seen.add(identity)
        median = as_int(row[median_col] if median_col < len(row) else 0)
        raw.append((row_number, label, median, link))
        if len(raw) >= limit:
            break

    forms_by_link: dict[str, set[str]] = defaultdict(set)
    forms_sheet_found = False
    for sheet in workbook.worksheets:
        if sheet is main:
            continue
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = header_map(first)
        form_col = find_column(headers, "Слова", "Слово")
        relation_col = find_column(headers, "Связь слов с другим листом")
        if form_col is None or relation_col is None:
            continue
        forms_sheet_found = True
        for row in sheet.iter_rows(min_row=2, values_only=True):
            form = normalize(row[form_col] if form_col < len(row) else "")
            relation = normalize(row[relation_col] if relation_col < len(row) else "")
            if form and relation:
                forms_by_link[relation].add(form)

    if not forms_sheet_found:
        warnings.append("Лист словоформ не найден; использованы только представительные формы")

    groups: list[Group] = []
    for row_number, label, median, link in raw:
        forms = set(forms_by_link.get(link, set())) if link else set()
        forms.add(label)
        groups.append(Group(row_number, label, median, tuple(sorted(forms)), link))

    form_owners: dict[str, list[str]] = defaultdict(list)
    for group in groups:
        for form in group.forms:
            form_owners[form].append(group.label)
    overlaps = {form: owners for form, owners in form_owners.items() if len(owners) > 1}
    if overlaps:
        sample = ", ".join(f"{form}: {owners}" for form, owners in list(overlaps.items())[:5])
        warnings.append(f"Есть формы в нескольких группах: {sample}")

    return groups, warnings


def read_terms(path: Path | None) -> set[str]:
    if path is None:
        return set()
    values: set[str] = set()
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        value = normalize(line.split("#", 1)[0])
        if value:
            values.add(value)
    return values


def group_selected(group: Group, include: set[str], exclude: set[str]) -> bool:
    return selection_status(group, include, exclude) == "eligible"


def selection_status(group: Group, include: set[str], exclude: set[str]) -> str:
    names = {group.label, *group.forms}
    if not names.isdisjoint(exclude):
        return "excluded"
    if include and names.isdisjoint(include):
        return "unclassified"
    return "eligible"


def count_group(text_tokens: list[str], group: Group) -> int:
    form_set = set(group.forms)
    return sum(1 for token in text_tokens if token in form_set)


def decision_integrity(groups: list[Group], include: set[str], exclude: set[str]) -> dict[str, int]:
    known = {name for group in groups for name in {group.label, *group.forms}}
    unclassified = 0
    conflicts = 0
    for group in groups:
        names = {group.label, *group.forms}
        in_include = not names.isdisjoint(include)
        in_exclude = not names.isdisjoint(exclude)
        unclassified += not in_include and not in_exclude
        conflicts += in_include and in_exclude
    return {
        "decision_unclassified": unclassified,
        "decision_conflicts": conflicts,
        "unknown_include_terms": len(include - known),
        "unknown_exclude_terms": len(exclude - known),
    }


def paragraph_count(text: str, group: Group) -> int:
    return sum(count_group(tokens(block), group) > 0 for block in PARAGRAPH_RE.split(text) if block.strip())


def coverage_and_depth(rows: list[dict[str, object]], side: str) -> tuple[int, int, float, float]:
    eligible = [row for row in rows if row["eligible"]]
    covered = sum(int(row[side]) > 0 for row in eligible)
    breadth = covered / len(eligible) if eligible else 0.0
    total_target = sum(int(row["target"]) for row in eligible)
    depth_numerator = sum(min(int(row[side]), int(row["target"])) for row in eligible)
    depth = depth_numerator / total_target if total_target else 0.0
    return covered, len(eligible), breadth, depth


def content_counter(text: str, target_forms: set[str]) -> Counter[str]:
    return Counter(
        token for token in tokens(text)
        if len(token) > 2 and token not in STOPWORDS and token not in target_forms
    )


def ngrams(text: str, size: int, target_forms: set[str]) -> Counter[str]:
    stream = tokens(text)
    result: Counter[str] = Counter()
    for index in range(len(stream) - size + 1):
        gram = stream[index:index + size]
        if all(token in STOPWORDS for token in gram):
            continue
        if any(token in target_forms for token in gram):
            continue
        result[" ".join(gram)] += 1
    return result


def anti_monoculture(source: str, edited: str, target_forms: set[str]) -> dict[str, list[dict[str, int | str]]]:
    before_words = content_counter(source, target_forms)
    after_words = content_counter(edited, target_forms)
    repeated_words = [
        {"text": word, "before": before_words[word], "after": count, "added": count - before_words[word]}
        for word, count in after_words.most_common()
        if count >= 3 and count - before_words[word] >= 2
    ]

    repeated_ngrams: list[dict[str, int | str]] = []
    for size in (2, 3):
        before = ngrams(source, size, target_forms)
        after = ngrams(edited, size, target_forms)
        repeated_ngrams.extend(
            {"text": gram, "before": before[gram], "after": count, "added": count - before[gram]}
            for gram, count in after.most_common()
            if count >= 2 and count - before[gram] >= 2
        )
    repeated_ngrams.sort(key=lambda item: (-int(item["added"]), -int(item["after"]), str(item["text"])))
    return {"words": repeated_words[:30], "ngrams": repeated_ngrams[:30]}


def local_repetitions(text: str, groups: Iterable[object]) -> dict[str, list[dict[str, object]]]:
    raw_blocks = [part for part in PARAGRAPH_RE.split(text) if part.strip()]
    blocks: list[str] = []
    for block in raw_blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if len(lines) > 1 or any(line.startswith(("#", "- ", "* ")) for line in lines):
            blocks.extend(line for line in lines if not line.startswith("#"))
        elif lines and not lines[0].startswith("#"):
            blocks.append(lines[0])

    sentences: list[str] = []
    sentence_ranges: list[tuple[int, int]] = []
    for block in blocks:
        start = len(sentences)
        sentences.extend(part.strip() for part in SENTENCE_RE.split(block) if part.strip())
        sentence_ranges.append((start, len(sentences)))
    sentence_tokens = [tokens(part) for part in sentences]
    same_sentence_words: list[dict[str, object]] = []
    adjacent_words: list[dict[str, object]] = []
    paragraph_words: list[dict[str, object]] = []

    for index, stream in enumerate(sentence_tokens, start=1):
        counts = Counter(token for token in stream if len(token) > 3 and token not in STOPWORDS)
        for word, count in counts.items():
            if count >= 2:
                same_sentence_words.append({"text": word, "sentence": index, "count": count})

    for start, end in sentence_ranges:
        for index in range(start, end - 1):
            left = Counter(token for token in sentence_tokens[index] if len(token) > 3 and token not in STOPWORDS)
            right = Counter(token for token in sentence_tokens[index + 1] if len(token) > 3 and token not in STOPWORDS)
            for word in sorted(left.keys() & right.keys()):
                adjacent_words.append({
                    "text": word,
                    "sentences": [index + 1, index + 2],
                    "count": left[word] + right[word],
                })

    same_sentence_groups: list[dict[str, object]] = []
    adjacent_groups: list[dict[str, object]] = []
    paragraph_groups: list[dict[str, object]] = []
    paragraphs = blocks
    for index, paragraph in enumerate(paragraphs, start=1):
        counts = Counter(token for token in tokens(paragraph) if len(token) > 3 and token not in STOPWORDS)
        for word, count in counts.items():
            if count >= 3:
                paragraph_words.append({"text": word, "paragraph": index, "count": count})
    for group in groups:
        label = str(getattr(group, "label"))
        sentence_counts = [count_group(stream, group) for stream in sentence_tokens]
        for index, count in enumerate(sentence_counts, start=1):
            if count >= 2:
                same_sentence_groups.append({"group": label, "sentence": index, "count": count})
        for start, end in sentence_ranges:
            for index in range(start, end - 1):
                if sentence_counts[index] and sentence_counts[index + 1]:
                    adjacent_groups.append({
                        "group": label,
                        "sentences": [index + 1, index + 2],
                        "count": sentence_counts[index] + sentence_counts[index + 1],
                    })
        for index, paragraph in enumerate(paragraphs, start=1):
            count = count_group(tokens(paragraph), group)
            if count >= 3:
                paragraph_groups.append({"group": label, "paragraph": index, "count": count})

    return {
        "same_sentence_words": same_sentence_words,
        "adjacent_words": adjacent_words,
        "paragraph_words": paragraph_words,
        "same_sentence_groups": same_sentence_groups,
        "adjacent_groups": adjacent_groups,
        "paragraph_groups": paragraph_groups,
    }


def audit(
    groups: list[Group],
    source: str,
    edited: str,
    include: set[str],
    exclude: set[str],
    core_size: int,
) -> dict[str, object]:
    before_tokens = tokens(source)
    after_tokens = tokens(edited)
    eligible_groups = [group for group in groups if group_selected(group, include, exclude)]
    ordered = sorted(eligible_groups, key=lambda group: group.row)
    core_labels = {group.label for group in ordered[:core_size]}

    rows: list[dict[str, object]] = []
    for group in groups:
        status = selection_status(group, include, exclude)
        eligible = status == "eligible"
        before = count_group(before_tokens, group)
        after = count_group(after_tokens, group)
        core = group.label in core_labels
        target = max(1, group.median) if core else (1 if eligible else 0)
        cap = target
        paragraphs = paragraph_count(edited, group)
        rows.append({
            "row": group.row,
            "group": group.label,
            "median": group.median,
            "target": target,
            "cap": cap,
            "before": before,
            "after": after,
            "added": after - before,
            "remaining": max(0, target - after),
            "eligible": eligible,
            "status": status,
            "core": core,
            "paragraphs": paragraphs,
            "over_median": eligible and after > cap,
            "clustered": eligible and after >= 2 and paragraphs < 2,
            "forms": list(group.forms),
        })

    covered_before, eligible_count, breadth_before, depth_before = coverage_and_depth(rows, "before")
    covered_after, _, breadth_after, depth_after = coverage_and_depth(rows, "after")
    target_forms = {form for group in eligible_groups for form in group.forms}
    monoculture = anti_monoculture(source, edited, target_forms)
    local = local_repetitions(edited, eligible_groups)

    return {
        "summary": {
            "groups_in_pool": len(groups),
            "eligible_groups": eligible_count,
            "excluded_groups": sum(row["status"] == "excluded" for row in rows),
            "unclassified_groups": sum(row["status"] == "unclassified" for row in rows),
            "zero_median_groups": sum(int(row["median"]) == 0 for row in rows),
            "covered_before": covered_before,
            "covered_after": covered_after,
            "breadth_before": round(breadth_before, 4),
            "breadth_after": round(breadth_after, 4),
            "depth_before": round(depth_before, 4),
            "depth_after": round(depth_after, 4),
            "over_median": sum(bool(row["over_median"]) and bool(row["eligible"]) for row in rows),
            "missing_width": sum(bool(row["eligible"]) and int(row["after"]) == 0 for row in rows),
            "target_deficits": sum(bool(row["eligible"]) and int(row["remaining"]) > 0 for row in rows),
            "clustered": sum(bool(row["clustered"]) for row in rows),
            "yo_symbols": edited.count("\u0451") + edited.count("\u0401"),
            "monoculture_words": len(monoculture["words"]),
            "monoculture_ngrams": len(monoculture["ngrams"]),
            "same_sentence_repetitions": len(local["same_sentence_groups"]),
            "adjacent_sentence_repetitions": len(local["adjacent_groups"]),
            "paragraph_repetitions": len(local["paragraph_groups"]),
            "same_sentence_word_repetitions": len(local["same_sentence_words"]),
            "adjacent_sentence_word_repetitions": len(local["adjacent_words"]),
            "paragraph_word_repetitions": len(local["paragraph_words"]),
        },
        "groups": rows,
        "anti_monoculture": monoculture,
        "local_repetitions": local,
    }


def render(report: dict[str, object], warnings: list[str]) -> None:
    summary = report["summary"]
    assert isinstance(summary, dict)
    print(
        "POOL={groups_in_pool} ELIGIBLE={eligible_groups} "
        "BREADTH={covered_before}->{covered_after} "
        "({breadth_before:.1%}->{breadth_after:.1%}) "
        "DEPTH={depth_before:.1%}->{depth_after:.1%}".format(**summary)
    )
    print(
        "OVER_MEDIAN={over_median} CLUSTERED={clustered} "
        "MONOCULTURE_WORDS={monoculture_words} "
        "MONOCULTURE_NGRAMS={monoculture_ngrams} YO={yo_symbols}".format(**summary)
    )
    print(
        "LOCAL_REPETITIONS sentence={same_sentence_repetitions} "
        "adjacent={adjacent_sentence_repetitions} paragraph={paragraph_repetitions} "
        "words={same_sentence_word_repetitions}/{adjacent_sentence_word_repetitions}/{paragraph_word_repetitions}".format(**summary)
    )
    print(
        "DECISIONS unclassified={decision_unclassified} conflicts={decision_conflicts} "
        "unknown={unknown_include_terms}/{unknown_exclude_terms} input_warnings={input_warnings}".format(**summary)
    )
    for warning in warnings:
        print(f"WARNING: {warning}")

    groups = report["groups"]
    assert isinstance(groups, list)
    for row in groups:
        if row["core"]:
            print(
                f"CORE: {row['group']} "
                f"{row['before']}->{row['after']}/{row['target']} paragraphs={row['paragraphs']}"
            )
    problems = [
        row for row in groups
        if row["eligible"] and (row["over_median"] or row["clustered"])
    ]
    for row in problems:
        flags = []
        if row["over_median"]:
            flags.append("OVER")
        if row["clustered"]:
            flags.append("CLUSTERED")
        print(
            f"{'+'.join(flags)}: {row['group']} "
            f"{row['before']}->{row['after']} target={row['target']} paragraphs={row['paragraphs']}"
        )

    missing = [row["group"] for row in groups if row["eligible"] and int(row["after"]) == 0]
    if missing:
        print("UNCOVERED: " + ", ".join(str(item) for item in missing))

    anti = report["anti_monoculture"]
    assert isinstance(anti, dict)
    for item in anti["words"][:10]:
        print(f"WRAPPER_WORD: {item['text']} {item['before']}->{item['after']}")
    for item in anti["ngrams"][:10]:
        print(f"WRAPPER_NGRAM: {item['text']} {item['before']}->{item['after']}")
    local = report["local_repetitions"]
    assert isinstance(local, dict)
    for item in local["same_sentence_groups"][:10]:
        print(f"SENTENCE_REPEAT: {item['group']} sentence={item['sentence']} count={item['count']}")
    for item in local["adjacent_groups"][:10]:
        print(f"ADJACENT_REPEAT: {item['group']} sentences={item['sentences']} count={item['count']}")


def strict_failures(report: dict[str, object]) -> list[str]:
    summary = report["summary"]
    assert isinstance(summary, dict)
    failures: list[str] = []
    if int(summary["missing_width"]):
        failures.append("есть непокрытые релевантные группы")
    if int(summary["target_deficits"]):
        failures.append("есть недобор целевых значений")
    if int(summary["over_median"]):
        failures.append("есть превышения медианы")
    if int(summary.get("input_warnings", 0)):
        failures.append("есть ошибки структуры LSI-таблицы")
    if int(summary.get("decision_unclassified", 0)):
        failures.append("не для всех групп записано решение")
    if int(summary.get("decision_conflicts", 0)):
        failures.append("группа одновременно включена и исключена")
    if int(summary.get("unknown_include_terms", 0)) or int(summary.get("unknown_exclude_terms", 0)):
        failures.append("файлы решений содержат неизвестные группы или формы")
    if int(summary["yo_symbols"]):
        failures.append("найден символ U+0451")
    if int(summary["clustered"]):
        failures.append("есть группы, собранные в одном абзаце")
    if int(summary["same_sentence_repetitions"]):
        failures.append("целевая группа повторяется в одном предложении")
    if int(summary["adjacent_sentence_repetitions"]):
        failures.append("целевая группа повторяется в соседних предложениях")
    if int(summary["paragraph_repetitions"]):
        failures.append("целевая группа чрезмерно сконцентрирована в одном абзаце")
    if int(summary["same_sentence_word_repetitions"]):
        failures.append("знаменательное слово повторяется в одном предложении")
    if int(summary["adjacent_sentence_word_repetitions"]):
        failures.append("знаменательное слово повторяется в соседних предложениях")
    if int(summary["paragraph_word_repetitions"]):
        failures.append("знаменательное слово чрезмерно сконцентрировано в одном абзаце")
    if int(summary["covered_before"]) < int(summary["eligible_groups"]) and int(summary["covered_after"]) <= int(summary["covered_before"]):
        failures.append("семантический охват не вырос")
    if float(summary["depth_after"]) < float(summary["depth_before"]):
        failures.append("семантическая глубина снизилась")
    return failures


def self_test() -> int:
    groups = [
        Group(2, "душевой", 4, ("душевая", "душевой"), "1"),
        Group(3, "смеситель", 3, ("смеситель", "смесителя"), "2"),
        Group(4, "керамика", 2, ("керамика", "керамики", "керамические"), "3"),
        Group(5, "монтаж", 2, ("монтаж", "монтажа"), "4"),
        Group(6, "комплект", 1, ("комплект",), "5"),
        Group(7, "конфиденциальность", 10, ("конфиденциальность",), "6"),
        Group(8, "пароль", 8, ("пароль",), "7"),
    ]
    source = "Товар удобный. Удобный товар.\n\nОбщий товар для дома."
    edited = (
        "Душевая зона включает комплект из керамики.\n\n"
        "Смеситель рассчитан на аккуратный монтаж.\n\n"
        "Душевой уголок дополняют защитные детали.\n\n"
        "Для душевой ниши выбран смеситель.\n\n"
        "Душевой блок и смеситель упрощают порядок установки."
    )
    exclude = {"конфиденциальность", "пароль"}
    report = audit(groups, source, edited, set(), exclude, core_size=2)
    local_bad = local_repetitions(
        "Смеситель согласуют со схемой, смеситель устанавливают после замеров. Смеситель подключают после отделки.",
        [groups[1]],
    )
    rows = {str(row["group"]): row for row in report["groups"]}
    expected = {"душевой": 4, "смеситель": 3, "керамика": 1, "монтаж": 1, "комплект": 1}
    checks = {
        "core_reaches_median": all(int(rows[name]["after"]) == value for name, value in expected.items()),
        "no_group_over_median": all(not bool(row["over_median"]) for row in rows.values()),
        "noise_stays_zero": int(rows["конфиденциальность"]["after"]) == 0 and int(rows["пароль"]["after"]) == 0,
        "core_has_priority": int(rows["душевой"]["after"]) > int(rows["керамика"]["after"]) and int(rows["смеситель"]["after"]) > int(rows["комплект"]["after"]),
        "core_is_distributed": int(rows["душевой"]["paragraphs"]) >= 2 and int(rows["смеситель"]["paragraphs"]) >= 2,
        "breadth_is_complete": report["summary"]["covered_after"] == report["summary"]["eligible_groups"],
        "no_wrapper_monoculture": not report["anti_monoculture"]["words"] and not report["anti_monoculture"]["ngrams"],
        "same_sentence_repeat_detected": len(local_bad["same_sentence_groups"]) == 1,
        "adjacent_sentence_repeat_detected": len(local_bad["adjacent_groups"]) == 1,
        "same_sentence_word_repeat_detected": len(local_bad["same_sentence_words"]) == 1,
        "adjacent_sentence_word_repeat_detected": len(local_bad["adjacent_words"]) == 1,
    }
    for name, passed in checks.items():
        print(f"{'PASS' if passed else 'FAIL'} {name}")
    return 0 if all(checks.values()) else 1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, help="Исходный текст UTF-8")
    parser.add_argument("--text", type=Path, help="Отредактированный текст UTF-8")
    parser.add_argument("--lsi", type=Path, help="Excel-таблица LSI")
    parser.add_argument("--limit", type=int, default=200, help="Число первых групп для смыслового отбора и внутреннего покрытия")
    parser.add_argument("--core-size", type=int, default=12, choices=range(10, 13), help="Число первых групп ядра: 10–12")
    parser.add_argument("--include", type=Path, help="Белый список релевантных групп или форм")
    parser.add_argument("--exclude", type=Path, help="Список нерелевантных групп или форм")
    parser.add_argument("--json-output", type=Path, help="Сохранить полный аудит JSON")
    parser.add_argument("--strict", action="store_true", help="Вернуть ненулевой код при жестких нарушениях")
    parser.add_argument("--self-test", action="store_true", help="Запустить встроенный контрольный тест")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.self_test:
        return self_test()
    missing = [name for name in ("source", "text", "lsi") if getattr(args, name) is None]
    if missing:
        raise SystemExit("Обязательные параметры: --source, --text, --lsi")
    if args.limit < 1 or args.core_size < 1:
        raise SystemExit("--limit и --core-size должны быть положительными")

    groups, warnings = read_groups(args.lsi, args.limit)
    if args.strict and (args.include is None or args.exclude is None):
        raise SystemExit("В строгом режиме обязательны оба файла решений: --include и --exclude")
    source = args.source.read_text(encoding="utf-8-sig")
    edited = args.text.read_text(encoding="utf-8-sig")
    include = read_terms(args.include)
    exclude = read_terms(args.exclude)
    report = audit(groups, source, edited, include, exclude, args.core_size)
    report["summary"]["input_warnings"] = len(warnings)
    report["summary"].update(decision_integrity(groups, include, exclude))
    render(report, warnings)

    if args.json_output:
        args.json_output.parent.mkdir(parents=True, exist_ok=True)
        args.json_output.write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    failures = strict_failures(report)
    if args.strict and failures:
        print("STRICT_FAIL: " + "; ".join(failures), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
