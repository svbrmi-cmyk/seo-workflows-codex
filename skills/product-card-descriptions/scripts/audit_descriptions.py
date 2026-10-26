#!/usr/bin/env python3
"""Audit a product-description column in an XLSX workbook."""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument(
        "--description-column", required=True,
        help="Excel letter, one-based number, or exact header text",
    )
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument(
        "--data-key-column",
        help="Header used to identify data rows; defaults to column A",
    )
    parser.add_argument("--expected-rows", type=int)
    parser.add_argument("--min-chars", type=int)
    parser.add_argument("--max-chars", type=int)
    parser.add_argument("--html", action="store_true", dest="require_html")
    parser.add_argument("--require-unique", action="store_true")
    parser.add_argument(
        "--banned", action="append", default=[],
        help="Case-insensitive literal; may be repeated",
    )
    parser.add_argument(
        "--banned-regex", action="append", default=[],
        help="Case-insensitive regular expression; may be repeated",
    )
    parser.add_argument(
        "--fact-column", action="append", default=[],
        help="Header whose non-empty value must occur in the description",
    )
    parser.add_argument("--category-column")
    parser.add_argument(
        "--keyword-map",
        help="JSON object or path: source category -> required phrase",
    )
    parser.add_argument("--color-column")
    parser.add_argument(
        "--color-map",
        help="JSON object or path: source color -> required human label",
    )
    parser.add_argument("--max-errors", type=int, default=50)
    return parser.parse_args()


def load_json_argument(value: str | None) -> dict[str, str]:
    if not value:
        return {}
    path = Path(value)
    raw = path.read_text(encoding="utf-8") if path.is_file() else value
    parsed = json.loads(raw)
    if not isinstance(parsed, dict) or not all(
        isinstance(key, str) and isinstance(item, str)
        for key, item in parsed.items()
    ):
        raise ValueError("Mapping must be a JSON object with string keys and values")
    return parsed


def visible_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = html.unescape(re.sub(r"<[^>]+>", " ", text))
    return re.sub(r"\s+", " ", text).strip()


def normalized_fact(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace(".", ",")).strip().lower()


def resolve_column(spec: str, headers: list[Any]) -> int:
    if spec in headers:
        index = headers.index(spec) + 1
    elif spec.isdigit():
        index = int(spec)
    elif re.fullmatch(r"[A-Za-z]{1,3}", spec):
        index = column_index_from_string(spec.upper())
    else:
        raise ValueError(f"Header not found: {spec!r}")
    if index < 1:
        raise ValueError(f"Invalid column: {spec!r}")
    return index


def main() -> int:
    args = parse_args()
    keyword_map = load_json_argument(args.keyword_map)
    color_map = load_json_argument(args.color_map)

    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=False)
    if args.sheet not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {args.sheet!r}")
    sheet = workbook[args.sheet]
    headers = [
        cell.value
        for cell in next(
            sheet.iter_rows(min_row=args.header_row, max_row=args.header_row)
        )
    ]

    description_col = resolve_column(args.description_column, headers)
    key_col = (
        resolve_column(args.data_key_column, headers)
        if args.data_key_column else 1
    )
    fact_cols = [
        (name, resolve_column(name, headers)) for name in args.fact_column
    ]
    category_col = (
        resolve_column(args.category_column, headers)
        if args.category_column else None
    )
    color_col = (
        resolve_column(args.color_column, headers)
        if args.color_column else None
    )

    errors: list[dict[str, Any]] = []
    descriptions: list[str] = []
    lengths: list[int] = []
    rows_checked = 0

    def add_error(row: int, code: str, detail: str) -> None:
        if len(errors) < args.max_errors:
            errors.append({"row": row, "code": code, "detail": detail})

    for row_number, values in enumerate(
        sheet.iter_rows(min_row=args.header_row + 1, values_only=True),
        start=args.header_row + 1,
    ):
        key_value = values[key_col - 1] if key_col <= len(values) else None
        if key_value is None or str(key_value).strip() == "":
            continue
        rows_checked += 1
        raw = values[description_col - 1] if description_col <= len(values) else None
        visible = visible_text(raw)
        descriptions.append(visible)
        lengths.append(len(visible))

        if not visible:
            add_error(row_number, "empty", "Description is empty")
            continue
        if args.min_chars is not None and len(visible) < args.min_chars:
            add_error(row_number, "too_short", f"{len(visible)} < {args.min_chars}")
        if args.max_chars is not None and len(visible) > args.max_chars:
            add_error(row_number, "too_long", f"{len(visible)} > {args.max_chars}")

        raw_text = "" if raw is None else str(raw)
        if args.require_html:
            if "<p>" not in raw_text or raw_text.count("<p>") != raw_text.count("</p>"):
                add_error(row_number, "html", "Missing or unbalanced <p> tags")
            if raw_text.count("<div") != raw_text.count("</div>"):
                add_error(row_number, "html", "Unbalanced <div> tags")

        lower = visible.lower()
        for phrase in args.banned:
            if phrase.lower() in lower:
                add_error(row_number, "banned", phrase)
        for pattern in args.banned_regex:
            if re.search(pattern, visible, flags=re.IGNORECASE):
                add_error(row_number, "banned_regex", pattern)

        for name, column in fact_cols:
            fact = normalized_fact(values[column - 1])
            if fact and fact not in lower.replace(".", ","):
                add_error(row_number, "missing_fact", f"{name}: {fact}")

        if category_col and keyword_map:
            category = "" if values[category_col - 1] is None else str(values[category_col - 1])
            required = keyword_map.get(category)
            if required is None:
                add_error(row_number, "unknown_category", category)
            elif required.lower() not in lower:
                add_error(row_number, "missing_keyword", required)

        if color_col and color_map:
            color = "" if values[color_col - 1] is None else str(values[color_col - 1])
            label = color_map.get(color)
            if label is None:
                add_error(row_number, "unknown_color", color)
            else:
                if color.lower() not in lower:
                    add_error(row_number, "missing_color_code", color)
                if label.lower() not in lower:
                    add_error(row_number, "missing_color_label", label)

    if args.expected_rows is not None and rows_checked != args.expected_rows:
        add_error(args.header_row, "row_count", f"{rows_checked} != {args.expected_rows}")

    duplicate_count = len(descriptions) - len(set(descriptions))
    if args.require_unique and duplicate_count:
        add_error(args.header_row, "duplicates", str(duplicate_count))

    report = {
        "workbook": str(args.workbook),
        "sheet": args.sheet,
        "rows_checked": rows_checked,
        "nonempty_descriptions": sum(bool(item) for item in descriptions),
        "unique_descriptions": len(set(descriptions)),
        "duplicate_count": duplicate_count,
        "min_visible_chars": min(lengths) if lengths else None,
        "max_visible_chars": max(lengths) if lengths else None,
        "error_count_shown": len(errors),
        "errors": errors,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if errors else 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print(json.dumps({"fatal_error": str(exc)}, ensure_ascii=False, indent=2))
        sys.exit(2)
